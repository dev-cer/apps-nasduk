import io
import os
import secrets
import sqlite3
from datetime import date, datetime, timedelta, timezone
from typing import List, Optional

from fastapi import FastAPI, Header, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel

app = FastAPI(title="Order Service - Nasi Uduk Mpo Rohaya")

DB_PATH = os.getenv("ORDERS_DB_PATH", "/data/orders.db")
ADMIN_TOKEN = os.getenv("ADMIN_TOKEN", "")
WIB = timezone(timedelta(hours=7))

# Nama paket yang tersedia di Menu Service
PACKAGES = {
    1: "Paket Komplit Arupa",
    2: "Paket Sederhana Mpo",
    3: "Paket Semur Mantap",
    4: "Paket Hemat Bihun Only",
}

class OrderRequest(BaseModel):
    package_id: int
    customer_name: str

class Order(BaseModel):
    order_id: int
    package_id: int
    package_name: str
    customer_name: str
    created_at: str

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

@app.on_event("startup")
def init_db():
    os.makedirs(os.path.dirname(DB_PATH) or ".", exist_ok=True)
    with get_db() as conn:
        conn.execute(
            """CREATE TABLE IF NOT EXISTS orders (
                order_id INTEGER PRIMARY KEY AUTOINCREMENT,
                package_id INTEGER NOT NULL,
                package_name TEXT NOT NULL,
                customer_name TEXT NOT NULL,
                created_at TEXT NOT NULL
            )"""
        )

def require_admin(token: Optional[str]):
    if not ADMIN_TOKEN:
        raise HTTPException(status_code=503, detail="ADMIN_TOKEN belum diatur di server")
    if not token or not secrets.compare_digest(token, ADMIN_TOKEN):
        raise HTTPException(status_code=401, detail="Token admin salah")

def fetch_orders(day: Optional[date]) -> List[Order]:
    query, params = "SELECT * FROM orders", ()
    if day:
        query, params = query + " WHERE substr(created_at, 1, 10) = ?", (day.isoformat(),)
    with get_db() as conn:
        rows = conn.execute(query + " ORDER BY order_id", params).fetchall()
    return [Order(**dict(r)) for r in rows]

@app.get("/healthz")
def healthz():
    return {"status": "ok"}

@app.get("/api/v1/orders", response_model=List[Order])
def list_orders(date: Optional[date] = None, x_admin_token: Optional[str] = Header(None)):
    require_admin(x_admin_token)
    return fetch_orders(date)

@app.get("/api/v1/orders/export.xlsx")
def export_orders(date: Optional[date] = None, x_admin_token: Optional[str] = Header(None)):
    require_admin(x_admin_token)
    orders = fetch_orders(date)

    wb = Workbook()
    ws = wb.active
    ws.title = "Pesanan"
    ws.append(["No. Pesanan", "Waktu (WIB)", "Nama Pemesan", "Paket ID", "Nama Paket"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B45309")
    for o in orders:
        waktu = datetime.fromisoformat(o.created_at).strftime("%Y-%m-%d %H:%M")
        ws.append([o.order_id, waktu, o.customer_name, o.package_id, o.package_name])

    # Rekap jumlah pesanan per paket
    rekap = wb.create_sheet("Rekap")
    rekap.append(["Nama Paket", "Jumlah"])
    for cell in rekap[1]:
        cell.font = Font(bold=True)
    for pid, name in PACKAGES.items():
        rekap.append([name, sum(1 for o in orders if o.package_id == pid)])
    rekap.append(["Total", len(orders)])
    rekap[rekap.max_row][0].font = Font(bold=True)

    for sheet, widths in ((ws, [12, 18, 28, 10, 26]), (rekap, [26, 10])):
        for col, width in zip("ABCDE", widths):
            sheet.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"pesanan-{date.isoformat() if date else 'semua'}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

@app.post("/api/v1/orders")
def create_order(req: OrderRequest):
    name = req.customer_name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nama pemesan wajib diisi")
    if req.package_id not in PACKAGES:
        raise HTTPException(status_code=404, detail=f"Paket ID {req.package_id} tidak ditemukan")

    with get_db() as conn:
        cur = conn.execute(
            "INSERT INTO orders (package_id, package_name, customer_name, created_at) VALUES (?, ?, ?, ?)",
            (req.package_id, PACKAGES[req.package_id], name, datetime.now(WIB).isoformat(timespec="seconds")),
        )
        order_id = cur.lastrowid
    return {"message": f"Pesanan {name} untuk paket {req.package_id} sudah diterima dapur", "order_id": order_id}
